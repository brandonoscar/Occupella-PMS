"""Exports: any report as CSV, Excel or PDF, and a PMC's whole ledger as CSV files in a zip.

A report is a function named trust_report_<name> (CLAUDE.md conventions). `run_report` looks
the name up in the catalog, so a caller can name nothing else, and casts each argument to the
type the function declares, so a day or an id can come in as text. It runs under the
connection's own role: Postgres decides who may run which report.

Every format keeps what the report says:
  - CSV (RFC 4180, UTF-8): numbers as written, days as YYYY-MM-DD, times in UTC (ISO 8601),
    NULL as an empty field;
  - Excel: numbers as numbers, days as dates, times as UTC to the second; text stays text,
    except that control characters XML can't hold read as "?" and line breaks as "\n";
  - PDF: a table on US Letter, landscape, its header row repeated on every page. Its built-in
    font covers Latin-1 only, so any other character prints as "?": CSV and Excel are the
    exact copies.
A spreadsheet runs text that starts with =, +, -, @, a tab or a carriage return as a formula.
CSV gives such text a leading apostrophe (OWASP's advice on CSV injection), and Excel stores
it as text. Names come from people outside the PMC (tenants, vendors), so this matters.

The ledger export is what a PMC takes with it on leaving: accounts.csv, transfers.csv and
entries.csv, every row of its ledger, nothing summarized.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import psycopg
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from psycopg import sql

REPORT_NAME = re.compile(r"trust_report_[a-z0-9_]+")
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
FORMATS = ("csv", "xlsx", "pdf")
# The zip's file times, fixed so the same ledger always gives the same bytes.
ZIP_TIME = (1980, 1, 1, 0, 0, 0)


class ExportError(ValueError):
    """The export was refused before anything ran: not a report, or the wrong arguments."""


class UnknownReport(ExportError):
    """No report by that name: not a trust_report_* function, or none that exists."""


@dataclass(frozen=True)
class Table:
    title: str
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]


def run_report(conn: psycopg.Connection[Any], report: str, args: Sequence[Any]) -> Table:
    """Run one report with its arguments and keep its rows as Postgres returned them."""
    if not REPORT_NAME.fullmatch(report):
        raise UnknownReport(f"{report!r} is not a report: reports are named trust_report_<name>")
    found = conn.execute(
        "SELECT p.proname, array(SELECT format_type(a.type, NULL)"
        " FROM unnest(p.proargtypes::oid[]) WITH ORDINALITY AS a (type, n) ORDER BY a.n)"
        " FROM pg_proc AS p WHERE p.pronamespace = 'public'::regnamespace"
        " AND p.proname = %s AND p.proretset",
        (report,),
    ).fetchall()
    if len(found) != 1:
        raise UnknownReport(f"no report named {report}")
    # The function's name and argument types come from the catalog, never from the caller.
    function: str = found[0][0]
    types: list[str] = found[0][1]
    if len(args) != len(types):
        raise ExportError(
            f"{report} takes {len(types)} arguments ({', '.join(types)}), not {len(args)}"
        )
    query = sql.SQL("SELECT * FROM {}({})").format(
        sql.Identifier(function),
        sql.SQL(", ").join(sql.SQL("%s::" + kind) for kind in types),
    )
    with conn.cursor() as cursor:
        cursor.execute(query, list(args))
        columns = tuple(column.name for column in cursor.description or [])
        rows = tuple(cursor.fetchall())
    title = report.removeprefix("trust_report_").replace("_", " ")
    return Table(f"{title} ({', '.join(str(arg) for arg in args)})", columns, rows)


def text(value: Any) -> str:
    """How a value reads in CSV and PDF."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.astimezone(UTC).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict | list):
        return json.dumps(value, sort_keys=True, default=str)
    return str(value)


def to_csv(table: Table) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(table.columns)
    for row in table.rows:
        writer.writerow(_csv_cell(value) for value in row)
    return out.getvalue().encode()


def _csv_cell(value: Any) -> str:
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return text(value)


def to_xlsx(table: Table, generated: datetime) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None  # a new workbook always has one sheet
    sheet.title = "report"
    book.properties.creator = "Occupella PMS"
    book.properties.title = _sheet_text(table.title)
    book.properties.created = book.properties.modified = _naive_utc(generated)
    sheet.append(list(table.columns))
    for heading in sheet[1]:
        heading.font = Font(bold=True)
    for r, row in enumerate(table.rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c)
            if isinstance(value, str):
                cell.value = _sheet_text(value)
                cell.data_type = "s"  # never a formula, whatever it starts with
            elif isinstance(value, datetime):
                # Excel keeps about a millisecond; the sheet shows and keeps whole seconds.
                cell.value = _naive_utc(value).replace(microsecond=0)
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(value, date):
                cell.value = value
                cell.number_format = "yyyy-mm-dd"
            elif value is None or isinstance(value, bool | int | float | Decimal):
                cell.value = value
            else:
                cell.value = _sheet_text(text(value))
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _sheet_text(value: str) -> str:
    """Text as a spreadsheet can hold it: control characters XML can't carry become "?", and
    line breaks are "\n" (XML would turn a carriage return into one when the file is read)."""
    return ILLEGAL_CHARACTERS_RE.sub("?", value).replace("\r\n", "\n").replace("\r", "\n")


def _naive_utc(moment: datetime) -> datetime:
    """Excel has no time zones: a time is stored as its UTC clock reading."""
    return moment.astimezone(UTC).replace(tzinfo=None)


class _Pdf(FPDF):
    def footer(self) -> None:
        self.set_y(-28)
        self.set_font("Helvetica", size=7)
        self.cell(0, 10, f"Page {self.page_no()} of {{nb}}", align="C")


def to_pdf(table: Table, generated: datetime) -> bytes:
    pdf = _Pdf(orientation="landscape", unit="pt", format="letter")
    pdf.set_creation_date(generated)
    pdf.set_title(table.title)
    pdf.set_author("Occupella PMS")
    pdf.set_margins(36, 36, 36)
    pdf.set_auto_page_break(True, margin=40)
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 16, _latin1(table.title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=7)
    stamp = generated.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")
    pdf.cell(0, 12, f"Generated {stamp}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    shares = _widths(table)
    widths = [pdf.epw * share / sum(shares) for share in shares]
    pdf.set_font("Helvetica", style="B", size=PDF_SIZE)
    heading = [
        _wrap(pdf, _latin1(column), width - 2 * PDF_PAD)
        for column, width in zip(table.columns, widths, strict=True)
    ]
    _pdf_row(pdf, heading, widths)
    pdf.set_font("Helvetica", size=PDF_SIZE)
    for row in table.rows:
        cells = [
            _wrap(pdf, _latin1(text(value)), width - 2 * PDF_PAD)
            for value, width in zip(row, widths, strict=True)
        ]
        if pdf.get_y() + _pdf_height(cells) > pdf.page_break_trigger:
            pdf.add_page()
            pdf.set_font("Helvetica", style="B", size=PDF_SIZE)
            _pdf_row(pdf, heading, widths)
            pdf.set_font("Helvetica", size=PDF_SIZE)
        _pdf_row(pdf, cells, widths)
    return bytes(pdf.output())


# A table's text: its size and line height, and the room left around it in each cell (points).
PDF_SIZE = 7
PDF_LINE = 9
PDF_PAD = 2


def _pdf_height(cells: list[list[str]]) -> float:
    return max(len(lines) for lines in cells) * PDF_LINE + PDF_PAD


def _pdf_row(pdf: FPDF, cells: list[list[str]], widths: list[float]) -> None:
    """One row of the table in the current font, each cell boxed with its lines of text. Drawn
    directly rather than through fpdf's tables, which took 17 s for a month's general ledger of
    a 1,000-unit company (tools/speed.py)."""
    height = _pdf_height(cells)
    x, y = pdf.l_margin, pdf.get_y()
    for lines, width in zip(cells, widths, strict=True):
        pdf.rect(x, y, width, height)
        for n, line in enumerate(lines):
            pdf.text(x + PDF_PAD, y + PDF_PAD + n * PDF_LINE + PDF_SIZE, line)
        x += width
    pdf.set_y(y + height)


def _wrap(pdf: FPDF, value: str, width: float) -> list[str]:
    """A cell's text in lines no wider than `width` points in the current font: broken at
    spaces where it can and inside a word where it must, its own line breaks kept."""
    font, size = pdf.current_font, pdf.font_size_pt
    assert font is not None  # to_pdf sets the font before any text

    def wide(words: str) -> float:
        # The font's own widths: fpdf's get_string_width copies its text fragments each call.
        return font.get_text_width(words, size, None)[1]

    space = wide(" ")
    lines = []
    for paragraph in value.split("\n"):
        line, room = "", 0.0
        for word in paragraph.split(" "):
            needs = wide(word)
            if not line and needs <= width or line and room + space + needs <= width:
                line, room = (f"{line} {word}", room + space + needs) if line else (word, needs)
                continue
            if line:
                lines.append(line)
            while needs > width:
                fits = 1
                while fits < len(word) and wide(word[: fits + 1]) <= width:
                    fits += 1
                lines.append(word[:fits])
                word = word[fits:]
                needs = wide(word)
            line, room = word, needs
        lines.append(line)
    return lines


def _latin1(value: str) -> str:
    """The PDF's built-in font has Latin-1 only; anything else prints as "?"."""
    return value.encode("latin-1", "replace").decode("latin-1")


def _widths(table: Table) -> tuple[int, ...]:
    """Each column as wide as its longest value (between 4 and 40 characters)."""
    longest = [len(column) for column in table.columns]
    for row in table.rows:
        for c, value in enumerate(row):
            longest[c] = max(longest[c], len(text(value)))
    return tuple(min(max(n, 4), 40) for n in longest)


def export(table: Table, fmt: str, generated: datetime | None = None) -> bytes:
    """The table in one of FORMATS."""
    moment = generated or datetime.now(UTC)
    if fmt == "csv":
        return to_csv(table)
    if fmt == "xlsx":
        return to_xlsx(table, moment)
    if fmt == "pdf":
        return to_pdf(table, moment)
    raise ExportError(f"no export format {fmt!r}; use one of {', '.join(FORMATS)}")


LEDGER_ACCOUNTS = """
SELECT * FROM (
SELECT t.ledger_account_id, b.display_name || ' (' || b.kind || ')' AS bank, t.kind,
       CASE t.kind
           WHEN 'bank_cash' THEN b.display_name
           WHEN 'owner_property' THEN o.display_name || ' / ' || pr.display_name
           WHEN 'vendor_payable' THEN v.display_name
           WHEN 'pmc_income' THEN pmc.display_name
           ELSE tn.display_name
       END AS holder,
       a.balance, a.created_at AS opened_at
FROM trust_ledger_accounts AS t
JOIN pgledger_accounts AS a ON a.id = t.ledger_account_id
JOIN trust_bank_accounts AS b ON b.id = t.bank_account_id
JOIN trust_pmcs AS pmc ON pmc.pmc_id = t.pmc_id
LEFT JOIN trust_owners AS o ON o.id = t.owner_id
LEFT JOIN trust_properties AS pr ON pr.id = t.property_id
LEFT JOIN trust_tenants AS tn ON tn.id = t.tenant_id
LEFT JOIN trust_vendors AS v ON v.id = t.vendor_id
WHERE t.pmc_id = %(pmc)s::uuid
) AS a
ORDER BY a.bank COLLATE "C", a.kind COLLATE "C", a.holder COLLATE "C", a.ledger_account_id
"""

# A transfer never crosses PMCs (CLAUDE.md), so its from-account places it.
LEDGER_TRANSFERS = """
SELECT tr.id AS transfer_id, tr.event_at AS posted_at, tr.created_at AS recorded_at,
       tr.from_account_id, tr.to_account_id, tr.amount, tr.metadata ->> 'memo' AS memo,
       tr.metadata
FROM pgledger_transfers AS tr
JOIN trust_ledger_accounts AS t ON t.ledger_account_id = tr.from_account_id
WHERE t.pmc_id = %(pmc)s::uuid
ORDER BY tr.event_at, tr.created_at, tr.id
"""

LEDGER_ENTRIES = """
SELECT e.id AS entry_id, e.transfer_id, e.account_id, e.amount,
       e.account_previous_balance AS balance_before, e.account_current_balance AS balance_after,
       e.account_version, e.created_at AS recorded_at
FROM pgledger_entries AS e
JOIN trust_ledger_accounts AS t ON t.ledger_account_id = e.account_id
WHERE t.pmc_id = %(pmc)s::uuid
ORDER BY e.account_id, e.account_version
"""


def ledger_zip(conn: psycopg.Connection[Any], pmc_id: Any) -> bytes:
    """Every account, transfer and entry of one PMC's ledger, as three CSV files."""
    known = conn.execute("SELECT 1 FROM trust_pmcs WHERE pmc_id = %s::uuid", (pmc_id,))
    if known.fetchone() is None:
        raise ExportError(f"no PMC {pmc_id}")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, query in (
            ("accounts", LEDGER_ACCOUNTS),
            ("transfers", LEDGER_TRANSFERS),
            ("entries", LEDGER_ENTRIES),
        ):
            with conn.cursor() as cursor:
                cursor.execute(query, {"pmc": pmc_id})
                columns = tuple(column.name for column in cursor.description or [])
                table = Table(name, columns, tuple(cursor.fetchall()))
            info = zipfile.ZipInfo(f"{name}.csv", ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, to_csv(table))
    return out.getvalue()


# Everything of one PMC (the founder's review q132): every record table as CSV, its files as
# files, and the ledger. Tables whose rows are credentials are never in it, nor the hash of any
# token (an invite's, a signing link's...), in a table or in the audit log's copies of rows.
NEVER_EXPORTED = frozenset({"trust_login_secrets", "trust_sessions"})
# An API key's secret_sha256 too (ADR 0011).
CREDENTIAL_COLUMNS = ("token_hash", "token_sha256", "status_token_sha256", "secret_sha256")
# The binary columns that are files (a document's bytes, a lease offered, a signature drawn):
# each written as a file of its own under files/, the CSV naming it. Other binary columns are
# digests, written as hex.
FILE_COLUMNS = frozenset(
    {
        ("trust_document_files", "content"),
        ("trust_lease_offers", "document"),
        ("trust_signature_evidence", "drawing"),
    }
)
RECORD_TABLES = """
SELECT c.relname
FROM pg_class AS c
JOIN pg_namespace AS n ON n.oid = c.relnamespace
WHERE n.nspname = 'public' AND c.relkind = 'r' AND c.relname LIKE 'trust\\_%%'
  AND EXISTS (
      SELECT 1 FROM pg_attribute AS a
      WHERE a.attrelid = c.oid AND a.attname = 'pmc_id' AND NOT a.attisdropped
  )
  AND has_table_privilege(c.oid, 'SELECT')
ORDER BY c.relname COLLATE "C"
"""
RECORD_COLUMNS = """
SELECT a.attname, format_type(a.atttypid, a.atttypmod)
FROM pg_attribute AS a
WHERE a.attrelid = %s::regclass AND a.attnum > 0 AND NOT a.attisdropped
ORDER BY a.attnum
"""


def _extension(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpg"
    return "bin"


def _add(archive: zipfile.ZipFile, name: str, data: bytes) -> None:
    info = zipfile.ZipInfo(name, ZIP_TIME)
    info.compress_type = zipfile.ZIP_DEFLATED
    archive.writestr(info, data)


def records_zip(conn: psycopg.Connection[Any], pmc_id: Any) -> bytes:
    """Every record of one PMC: records/<table>.csv for each of its tables the caller may read
    (rows in a fixed order, so the same books give the same file), files/<table>/... for its
    documents, leases offered and signatures drawn, and ledger/ as ledger_zip has it."""
    known = conn.execute("SELECT 1 FROM trust_pmcs WHERE pmc_id = %s::uuid", (pmc_id,))
    if known.fetchone() is None:
        raise ExportError(f"no PMC {pmc_id}")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        tables = [t for (t,) in conn.execute(RECORD_TABLES).fetchall() if t not in NEVER_EXPORTED]
        for table in tables:
            columns = [
                (name, kind)
                for name, kind in conn.execute(RECORD_COLUMNS, (table,)).fetchall()
                if name not in CREDENTIAL_COLUMNS
            ]
            picked = [
                sql.SQL("{} - {}::text[]").format(
                    sql.Identifier(name), sql.Literal(list(CREDENTIAL_COLUMNS))
                )
                if kind == "jsonb"
                else sql.Identifier(name)
                for name, kind in columns
            ]
            where: sql.Composable = sql.SQL("pmc_id = %(pmc)s::uuid")
            if table == "trust_audit_log":
                where = sql.SQL("{} AND table_name <> ALL(%(never)s)").format(where)
            query = sql.SQL(
                'SELECT {} FROM {} AS r WHERE {} ORDER BY row_to_json(r)::text COLLATE "C"'
            ).format(sql.SQL(", ").join(picked), sql.Identifier(table), where)
            rows = conn.execute(query, {"pmc": pmc_id, "never": list(NEVER_EXPORTED)}).fetchall()
            written: list[tuple[Any, ...]] = []
            for n, row in enumerate(rows, start=1):
                cells: list[Any] = []
                for (name, _), value in zip(columns, row, strict=True):
                    if isinstance(value, bytes | memoryview):
                        data = bytes(value)
                        if (table, name) in FILE_COLUMNS:
                            path = f"files/{table}/{n:06d}-{name}.{_extension(data)}"
                            _add(archive, path, data)
                            value = path
                        else:
                            value = data.hex()
                    cells.append(value)
                written.append(tuple(cells))
            names = tuple(name for name, _ in columns)
            _add(archive, f"records/{table}.csv", to_csv(Table(table, names, tuple(written))))
        for name, of_ledger in (
            ("accounts", LEDGER_ACCOUNTS),
            ("transfers", LEDGER_TRANSFERS),
            ("entries", LEDGER_ENTRIES),
        ):
            with conn.cursor() as cursor:
                cursor.execute(of_ledger, {"pmc": pmc_id})
                columns_of = tuple(column.name for column in cursor.description or [])
                ledger = Table(name, columns_of, tuple(cursor.fetchall()))
            _add(archive, f"ledger/{name}.csv", to_csv(ledger))
    return out.getvalue()


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m pms.exports",
        description="Export a report (CSV, Excel or PDF) or a PMC's whole ledger (a zip of CSV)."
        " Connects with DATABASE_URL.",
    )
    parser.add_argument("--role", help="run as this database role, e.g. trust_app")
    parser.add_argument("--out", help="write here instead of standard output")
    commands = parser.add_subparsers(dest="command", required=True)
    report = commands.add_parser("report", help="one report")
    report.add_argument("name", help="the report function, e.g. trust_report_rent_roll")
    report.add_argument("args", nargs="*", help="its arguments, in order")
    report.add_argument("--format", choices=FORMATS, default="csv")
    ledger = commands.add_parser("ledger", help="a PMC's accounts, transfers and entries")
    ledger.add_argument("pmc_id")
    options = parser.parse_args(argv)

    with psycopg.connect(os.environ["DATABASE_URL"]) as conn:
        if options.role:
            conn.execute(sql.SQL("SET ROLE {}").format(sql.Identifier(options.role)))
        try:
            if options.command == "report":
                data = export(run_report(conn, options.name, options.args), options.format)
            else:
                data = ledger_zip(conn, options.pmc_id)
        except ExportError as refused:
            print(f"refused: {refused}", file=sys.stderr)
            return 2
    if options.out:
        with open(options.out, "wb") as target:
            target.write(data)
    else:
        sys.stdout.buffer.write(data)
    return 0


if __name__ == "__main__":
    sys.exit(main())
